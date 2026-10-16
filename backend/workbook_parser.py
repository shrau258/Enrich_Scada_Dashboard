from __future__ import annotations
from io import BytesIO
import re
from typing import Any
from openpyxl import load_workbook

def parse_workbook(data:bytes,filename:str)->dict[str,Any]:
    workbook=load_workbook(BytesIO(data),read_only=True,data_only=True)
    sheet=workbook.active
    def text(row:int,col:int)->str: return str(sheet.cell(row,col).value or '').strip()
    customer=text(2,2).split(':',1)[-1].strip()
    plant=text(3,2).split(':',1)[-1].strip().rstrip('.')
    capacities=[float(x) for x in re.findall(r'\d+(?:\.\d+)?',text(4,2))]
    inverter_count=int(re.search(r'\d+',text(5,4)).group())
    inverter=[]; devices=[]; group=None
    for row in range(16,sheet.max_row+1):
        inv_name=sheet.cell(row,3).value; inv_address=sheet.cell(row,4).value
        if inv_name and isinstance(inv_address,(int,float)): inverter.append({'parameter':str(inv_name).strip(),'address':int(inv_address)})
        device_name=sheet.cell(row,8).value; device_address=sheet.cell(row,9).value
        if sheet.cell(row,10).value: group=str(sheet.cell(row,10).value).strip()
        if device_name and isinstance(device_address,(int,float)): devices.append({'group':group or 'Ungrouped','parameter':str(device_name).strip(),'address':int(device_address)})
    return {'filename':filename,'sheet':sheet.title,'customerName':customer,'plantName':plant,'acCapacityMw':capacities[0] if capacities else None,'dcCapacityMwp':capacities[1] if len(capacities)>1 else None,'inverterManufacturer':text(5,3),'inverterCount':inverter_count,'modbusFunctionCode':text(6,3),'sharedIp':text(9,3),'basePort':int(sheet.cell(10,3).value),'deviceArrayPort':int(sheet.cell(10,8).value),'startingAddress':int(sheet.cell(11,3).value),'dataCount':int(sheet.cell(12,3).value),'inverterMappings':inverter,'deviceArrayMappings':devices,'deviceArrayGroups':sorted({x['group'] for x in devices}),'warnings':['Register scaling, word order, signedness and decoder types are not defined in the workbook. Live Modbus activation remains blocked.']}
